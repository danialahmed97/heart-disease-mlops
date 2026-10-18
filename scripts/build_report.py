"""Convert docs/REPORT.md into docs/REPORT.docx and docs/REPORT.pdf.

A small Markdown subset is parsed (headings, paragraphs, lists, tables, code
fences, images, bold/inline-code) and rendered with python-docx and
reportlab. ``[SCREENSHOT: name.png]`` lines are replaced by
``screenshots/name.png`` when the file exists, otherwise by a grey note.
Mermaid blocks are replaced by a generated architecture diagram.
"""
import logging
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402
from PIL import Image as PILImage  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.enums import TA_CENTER  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import cm  # noqa: E402
from reportlab.platypus import (Image, ListFlowable, ListItem,  # noqa: E402
                                PageBreak, Paragraph, Preformatted,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SOURCE = DOCS / "REPORT.md"
SHOTS = ROOT / "screenshots"
ARCH_PNG = ROOT / "reports" / "figures" / "architecture.png"

logger = logging.getLogger("build_report")


@dataclass
class Block:
    """One parsed Markdown block."""

    kind: str  # title|h|p|ul|ol|table|code|img|shot|pagebreak
    text: str = ""
    level: int = 0
    items: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    path: Path | None = None


def make_architecture_png(path: Path = ARCH_PNG) -> Path:
    """Draw the pipeline architecture used in place of the Mermaid block."""
    path.parent.mkdir(parents=True, exist_ok=True)
    nodes = ["UCI data", "Download\n+ clean", "Train +\nGridSearchCV",
             "model.joblib\n+ metadata", "Docker\nimage", "Kubernetes\nDeploy+Svc+HPA",
             "Prometheus", "Grafana"]
    fig, ax = plt.subplots(figsize=(11, 3.4))
    ax.axis("off")
    xs = [0.06 + i * 0.125 for i in range(len(nodes))]
    for x, name in zip(xs, nodes):
        ax.text(x, 0.55, name, ha="center", va="center", fontsize=8,
                bbox=dict(boxstyle="round,pad=0.5", fc="#e8f0fe", ec="#3b6fd4"))
    for a, b in zip(xs, xs[1:]):
        ax.annotate("", xy=(b - 0.052, 0.55), xytext=(a + 0.052, 0.55),
                    arrowprops=dict(arrowstyle="->", color="#444"))
    ax.text(xs[2], 0.18, "MLflow ./mlruns", ha="center", fontsize=8,
            bbox=dict(boxstyle="round", fc="#fff4d6", ec="#c49a1a"))
    ax.annotate("", xy=(xs[2], 0.27), xytext=(xs[2], 0.42),
                arrowprops=dict(arrowstyle="->", color="#444"))
    ax.text(xs[4], 0.9, "GitHub Actions: lint > test > train > docker",
            ha="center", fontsize=8, style="italic")
    ax.annotate("", xy=(xs[4], 0.7), xytext=(xs[4], 0.85),
                arrowprops=dict(arrowstyle="->", ls="--", color="#888"))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return path


def parse(md: str) -> list[Block]:
    """Parse the Markdown subset into blocks."""
    lines = md.splitlines()
    blocks: list[Block] = []
    i = 0
    seen_title = False
    para: list[str] = []

    def flush() -> None:
        if para:
            blocks.append(Block("p", " ".join(s.strip() for s in para)))
            para.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            flush()
        elif s == "<!-- pagebreak -->":
            flush()
            blocks.append(Block("pagebreak"))
        elif s.startswith("```"):
            flush()
            lang, code = s[3:].strip(), []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            if lang == "mermaid":
                blocks.append(Block("img", "Architecture", path=make_architecture_png()))
            else:
                blocks.append(Block("code", "\n".join(code)))
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            text = s[level:].strip()
            if level == 1 and not seen_title:
                blocks.append(Block("title", text))
                seen_title = True
            else:
                blocks.append(Block("h", text, level=level))
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            i -= 1
            blocks.append(Block("table", rows=rows))
        elif re.match(r"^[*-] ", s) or re.match(r"^\d+\. ", s):
            flush()
            ordered = bool(re.match(r"^\d+\. ", s))
            items: list[str] = []
            while i < len(lines):
                t = lines[i]
                ts = t.strip()
                if re.match(r"^[*-] ", ts) and not t.startswith(" ") or \
                        re.match(r"^\d+\. ", ts):
                    items.append(re.sub(r"^([*-]|\d+\.) ", "", ts))
                elif t.startswith("  ") and ts:
                    sub = re.sub(r"^([*-]|\d+\.) ", "", ts)
                    if re.match(r"^[*-] ", ts):
                        items.append("- " + sub)  # flattened nested bullet
                    else:
                        items[-1] += " " + ts
                else:
                    break
                i += 1
            i -= 1
            blocks.append(Block("ol" if ordered else "ul", items=items))
        elif (m := re.match(r"^!\[(.*?)\]\((.*?)\)$", s)):
            flush()
            blocks.append(Block("img", m.group(1), path=(SOURCE.parent / m.group(2)).resolve()))
        elif (m := re.match(r"^\[SCREENSHOT: (.+?)\]$", s)):
            flush()
            blocks.append(Block("shot", m.group(1), path=SHOTS / m.group(1)))
        else:
            para.append(ln)
        i += 1
    flush()
    return blocks


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def inline_runs(text: str) -> list[tuple[str, str]]:
    """Split text into (segment, style) with style in {'', 'b', 'c'}."""
    out = []
    for part in INLINE.split(text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append((part[2:-2], "b"))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append((part[1:-1], "c"))
        elif part:
            out.append((part, ""))
    return out


def img_width_in(path: Path, max_in: float = 6.0) -> float:
    """Pick a display width: wide figures fill the page, small ones shrink."""
    with PILImage.open(path) as im:
        w, h = im.size
    return max_in if w / h > 1.35 else min(max_in, 4.0)


# ---------------------------------------------------------------- DOCX --
def build_docx(blocks: list[Block], out: Path) -> None:
    """Render blocks to a Word document."""
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    def add_runs(par, text: str) -> None:
        for seg, style in inline_runs(text):
            run = par.add_run(seg)
            if style == "b":
                run.bold = True
            elif style == "c":
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)

    def add_image(path: Path, caption: str) -> None:
        doc.add_picture(str(path), width=Inches(img_width_in(path)))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)

    on_title = False
    for b in blocks:
        if b.kind == "pagebreak":
            on_title = False
        if b.kind == "title":
            on_title = True
            for _ in range(5):
                doc.add_paragraph()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(b.text)
            r.bold = True
            r.font.size = Pt(26)
        elif b.kind == "h":
            doc.add_heading(b.text, level=min(b.level - 1, 3) or 1)
        elif b.kind == "p":
            par = doc.add_paragraph()
            if on_title:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(par, b.text)
        elif b.kind in ("ul", "ol"):
            style = "List Bullet" if b.kind == "ul" else "List Number"
            for it in b.items:
                nested = it.startswith("- ")
                par = doc.add_paragraph(style="List Bullet 2" if nested else style)
                add_runs(par, it[2:] if nested else it)
        elif b.kind == "table":
            ncols = max(len(r) for r in b.rows)
            t = doc.add_table(rows=len(b.rows), cols=ncols)
            t.style = "Light Grid Accent 1"
            for ri, row in enumerate(b.rows):
                for ci in range(ncols):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    add_runs(par, row[ci] if ci < len(row) else "")
                    for run in par.runs:
                        run.font.size = Pt(8.5)
                        if ri == 0:
                            run.bold = True
            doc.add_paragraph()
        elif b.kind == "code":
            par = doc.add_paragraph()
            run = par.add_run(b.text)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        elif b.kind == "img":
            if b.path and b.path.exists():
                add_image(b.path, b.text)
            else:
                logger.warning("Missing figure %s", b.path)
        elif b.kind == "shot":
            if b.path.exists():
                add_image(b.path, f"Screenshot: {b.text}")
            else:
                par = doc.add_paragraph()
                run = par.add_run(f"[Screenshot pending: {b.text}]")
                run.italic = True
                run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        elif b.kind == "pagebreak":
            doc.add_page_break()
    doc.save(out)


# ----------------------------------------------------------------- PDF --
def rl_markup(text: str) -> str:
    """Convert inline Markdown to reportlab mini-HTML."""
    out = []
    for seg, style in inline_runs(text):
        seg = escape(seg)
        if style == "b":
            out.append(f"<b>{seg}</b>")
        elif style == "c":
            out.append(f'<font name="Courier" size="8.5">{seg}</font>')
        else:
            out.append(seg)
    return "".join(out)


def build_pdf(blocks: list[Block], out: Path) -> None:
    """Render blocks to a PDF."""
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontSize=9.5, leading=13)
    cell = ParagraphStyle("cell", parent=body, fontSize=7.5, leading=9.5)
    cellb = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
    cap = ParagraphStyle("cap", parent=body, fontSize=8, alignment=TA_CENTER,
                         textColor=colors.grey, fontName="Helvetica-Oblique")
    title = ParagraphStyle("title", parent=ss["Title"], fontSize=26, leading=32)
    tcenter = ParagraphStyle("tc", parent=body, alignment=TA_CENTER, fontSize=12, leading=18)
    heads = {1: ss["Heading1"], 2: ss["Heading2"], 3: ss["Heading3"]}
    code = ParagraphStyle("code", fontName="Courier", fontSize=7.5, leading=9)
    story: list = []
    on_title = False

    def img(path: Path, caption: str) -> list:
        w_in = img_width_in(path)
        with PILImage.open(path) as im:
            ratio = im.size[1] / im.size[0]
        w = min(w_in * 72, 15.5 * cm)
        return [Image(str(path), width=w, height=w * ratio), Paragraph(escape(caption), cap),
                Spacer(1, 6)]

    for b in blocks:
        if b.kind == "title":
            story += [Spacer(1, 5 * cm), Paragraph(escape(b.text), title), Spacer(1, 1.2 * cm)]
            on_title = True
        elif b.kind == "pagebreak":
            story.append(PageBreak())
            on_title = False
        elif b.kind == "h":
            story.append(Paragraph(escape(b.text), heads[min(b.level - 1, 3) or 1]))
        elif b.kind == "p":
            story.append(Paragraph(rl_markup(b.text), tcenter if on_title else body))
            story.append(Spacer(1, 4))
        elif b.kind in ("ul", "ol"):
            items = [ListItem(Paragraph(rl_markup(i[2:] if i.startswith("- ") else i), body),
                              leftIndent=14 if i.startswith("- ") else 0)
                     for i in b.items]
            story.append(ListFlowable(items, bulletType="bullet" if b.kind == "ul" else "1",
                                      leftIndent=14, bulletFontSize=7))
            story.append(Spacer(1, 4))
        elif b.kind == "table":
            ncols = max(len(r) for r in b.rows)
            data = [[Paragraph(rl_markup(r[c] if c < len(r) else ""), cellb if ri == 0 else cell)
                     for c in range(ncols)] for ri, r in enumerate(b.rows)]
            t = Table(data, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dfe8f7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
            story += [t, Spacer(1, 8)]
        elif b.kind == "code":
            story += [Preformatted(b.text, code), Spacer(1, 6)]
        elif b.kind == "img":
            if b.path and b.path.exists():
                story += img(b.path, b.text)
        elif b.kind == "shot":
            if b.path.exists():
                story += img(b.path, f"Screenshot: {b.text}")
            else:
                story += [Paragraph(f"[Screenshot pending: {escape(b.text)}]", cap)]

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        if doc_.page > 1:
            canvas.drawCentredString(A4[0] / 2, 1.1 * cm, f"Page {doc_.page}")
        canvas.restoreState()

    SimpleDocTemplate(str(out), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm,
                      title="Heart Disease Prediction - MLOps Report",
                      author="Danial Barbhuiya").build(story, onFirstPage=footer,
                                                       onLaterPages=footer)


def main() -> int:
    """Build both output formats."""
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    blocks = parse(SOURCE.read_text(encoding="utf-8"))
    shots = [b for b in blocks if b.kind == "shot"]
    missing = [b.text for b in shots if not b.path.exists()]
    logger.info("%d blocks, %d screenshots referenced, %d missing",
                len(blocks), len(shots), len(missing))
    for name in missing:
        logger.info("  missing screenshots/%s (placeholder used)", name)
    build_docx(blocks, DOCS / "REPORT.docx")
    build_pdf(blocks, DOCS / "REPORT.pdf")
    logger.info("Wrote %s and %s", DOCS / "REPORT.docx", DOCS / "REPORT.pdf")
    return 0


if __name__ == "__main__":
    sys.exit(main())
